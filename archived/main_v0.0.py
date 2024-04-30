import streamlit as st
from docx import Document
from io import BytesIO

def load_docx(document):
    doc = Document(document)
    return doc

def modify_docx(doc, new_text):
    for para in doc.paragraphs:
        if '[targeted_text]' in para.text:
            para.text = new_text
    return doc

def save_docx(doc):
    buffer = BytesIO()
    doc.save(buffer)
    return buffer

def main():
    st.title('DOCX Editor App')

    uploaded_file = st.file_uploader("Choose a .docx file", type="docx")
    if uploaded_file is not None:
        document = load_docx(uploaded_file)
        print(document)
        user_input = st.text_input("New text to replace placeholders:", "Type here...")
        
        if st.button("Modify Document"):
            modified_doc = modify_docx(document, user_input)
            buffer = save_docx(modified_doc)
            buffer.seek(0)
            
            st.download_button(
                label="Download modified DOCX",
                data=buffer,
                file_name="modified_document.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

if __name__ == "__main__":
    main()